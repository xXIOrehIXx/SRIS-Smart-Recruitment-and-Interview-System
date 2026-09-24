"""
============================================================================
 BƯỚC 4 — GỘP MỌI SỐ ĐO VÀO MỘT FILE EXCEL

 Đọc: out/may_cham_4_ban.csv, out/<ver>/may_cham.csv (tầng máy)
      out/nguoi_cham_4_ban.csv, out/<ver>/nguoi_cham_tong_ket.csv,
      nguoi_cham_tung_cau.csv, nguoi_cham_bo_sot.csv (tầng người)
 -> out/KET_QUA_TONG_HOP.xlsx  +  out/<ver>/NGUOI_CHAM_<ver>.xlsx

 Sáu tab (cùng khuôn với exp_criteria_extract):
   1. DocTruoc          — ai chấm cái gì, file nào ra số nào, nguồn nhãn
   2. TongHop           — mỗi bậc một dòng, tầng người và tầng máy cạnh nhau
   3. TheoTin           — 10 CV x 4 bậc, xem CV nào kéo tụt chỉ số
   4. NguoiCham_TungCau — phiếu chấm tay, mỗi câu một dòng
   5. NguoiCham_BoSot   — mốc bị bỏ sót (mẫu số của recall)
   6. NguoiCham_TongKet — P/R/F1 từng CV + dòng TỔNG của mỗi bậc

 Tầng người chưa chấm thì các ô đó để trống — file vẫn dựng được.

 Chạy:  ..\\..\\ai-service\\.venv\\Scripts\\python.exe 4_gop_ket_qua_excel.py
============================================================================
"""

import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = Path(__file__).parent
OUT = HERE / "out"
VERS = ["v1", "v2", "v3", "v4"]

MO_TA_VER = {
    "v1": "Bảo tóm tắt, hết — không luật, không ép JSON",
    "v2": "+ Ràng buộc JSON schema (Pydantic) & temperature=0",
    "v3": "+ Luật chống bịa (chỉ dùng chữ trong CV, không suy diễn)",
    "v4": "+ Luật riêng của bản tóm tắt (3-5 câu, cấm tự cộng số năm, cấm chung chung) = PROMPT THẬT",
}

TEN_CV = {
    "C01_khong_ghi_so_nam": "Backend .NET — CV chỉ có mốc, không ghi số năm",
    "C02_ketoan_khop": "Kế toán — ca chuẩn, khớp yêu cầu",
    "C03_fresher_mong": "Frontend fresher — CV mỏng",
    "C04_trai_nganh": "Dân kinh doanh nộp vị trí IT",
    "C05_nhay_viec_nhieu_moc": "CSKH — 5 nơi làm trong 7 năm",
    "C06_text_pdf_lon_xon": "Vận hành CNC — text PDF dính chữ",
    "C07_song_ngu": "Marketing — CV song ngữ",
    "C08_thanh_tich_nhieu_so": "Trưởng nhóm KD — dày đặc con số",
    "C09_khoang_trong_su_nghiep": "Nhân sự C&B — nghỉ 2 năm giữa chừng",
    "C10_liet_ke_cong_nghe": "Backend — liệt kê cả rừng công nghệ",
}

NHAN_MO_TA = {
    "DUNG": "kể lại đúng một thông tin có thật trong CV",
    "BIA": "bịa — thông tin không hề có trong CV",
    "SUYDIEN": "suy diễn — có neo trong CV nhưng model suy ra thêm",
    "SAISO": "sai số — con số / mốc thời gian sai hoặc tự tính ra",
    "SAORONG": "sáo rỗng — gắn sang CV khác vẫn đúng",
}

NGUON_NHAN = ("Nhãn tầng người do trợ lý AI (Claude) soạn sơ bộ theo LUAT_NGUOI_CHAM.md ngày 25/09/2026; "
              "người làm đề tài duyệt lại, ưu tiên các dòng đánh dấu PHÂN VÂN. Phải nói rõ điều này khi trích số.")

XANH = PatternFill("solid", fgColor="1F4E79")
XAM = PatternFill("solid", fgColor="DDEBF7")
VANG = PatternFill("solid", fgColor="FFF2CC")
LUC = PatternFill("solid", fgColor="E2EFDA")
DO = PatternFill("solid", fgColor="FCE4E4")
TRANG = Font(color="FFFFFF", bold=True)
VIEN = Border(*[Side("thin", color="B0B0B0")] * 4)


def doc(path):
    p = Path(path)
    if not p.exists():
        return []
    with p.open(encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def so(v, mac_dinh=None):
    try:
        return float(v)
    except (TypeError, ValueError):
        return mac_dinh


def muc(v):
    if v is None:
        return ""
    return "Tốt" if v >= 0.85 else "Chấp nhận được" if v >= 0.70 else "Cần cải thiện"


def to_mau(v):
    return LUC if v >= 0.85 else DO if v < 0.70 else VANG


def dat_header(ws, hang, tieu_de):
    for i, t in enumerate(tieu_de, start=1):
        if t == "":          # ô đã bị merge từ hàng trên — ghi vào là AttributeError
            continue
        c = ws.cell(row=hang, column=i, value=t)
        c.fill, c.font = XANH, TRANG
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = VIEN


def rong(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def tieu_de_to(ws, hang, text, span):
    c = ws.cell(row=hang, column=1, value=text)
    c.font = Font(bold=True, size=12, color="1F4E79")
    ws.merge_cells(start_row=hang, start_column=1, end_row=hang, end_column=span)


def o(ws, hang, cot, gia_tri, **kw):
    c = ws.cell(row=hang, column=cot, value=gia_tri)
    c.border = VIEN
    c.alignment = Alignment(vertical="center", wrap_text=kw.get("wrap", False),
                            horizontal=kw.get("ngang", "left"))
    if "fmt" in kw:
        c.number_format = kw["fmt"]
    return c


# ---------------------------------------------------------------- tab 1
def tab_doc_truoc(wb):
    ws = wb.create_sheet("DocTruoc")
    rong(ws, [4, 34, 12, 70])

    tieu_de_to(ws, 1, "ĐỌC TRƯỚC — thí nghiệm này đo cái gì, ai chấm, file nào ra số nào", 4)
    ws.cell(row=2, column=2,
            value="Phần TÓM TẮT của lượt sàng lọc CV (/screen-cv) · 10 CV × 4 bậc prompt · qwen3:8b, temperature=0")
    ws.cell(row=3, column=2,
            value="Thí nghiệm BÓC LỚP (ablation): V4 là prompt đang chạy thật, V1–V3 là bản rút gọn của nó, CHƯA TỪNG chạy trong sản phẩm.")

    h = 5
    ws.cell(row=h, column=2, value="HAI TẦNG ĐO — vì sao cần cả hai").font = Font(bold=True, size=11, color="1F4E79")
    h += 1
    dat_header(ws, h, ["", "Tầng", "Ai chấm", "Trả lời câu hỏi gì · điểm mù"])
    for tang, ai, mo in [
        ("Tầng 1 — máy đo", "🤖 máy (may_cham.py)",
         "Đếm được: JSON hợp lệ, số câu, SỐ LẠ (con số không có trong CV), sáo rỗng, ổn định, giây/CV. "
         "ĐIỂM MÙ: không biết câu tóm tắt có ĐÚNG không — số tự cộng mà trùng một con số khác trong CV thì máy cho qua."),
        ("Tầng 2 — người chấm", "🧑 người (LUAT_NGUOI_CHAM.md)",
         "Gán 1 trong 5 nhãn cho TỪNG CÂU + đếm mốc bỏ sót → Precision / Recall / F1. "
         "ĐIỂM MÙ: không lặp lại được, phụ thuộc người chấm (xem số dòng PHÂN VÂN)."),
    ]:
        h += 1
        for i, g in enumerate(["", tang, ai, mo], start=1):
            if g:
                c = o(ws, h, i, g, wrap=True)
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if i == 2:
                    c.font = Font(bold=True)
        ws.row_dimensions[h].height = 62

    h += 2
    ws.cell(row=h, column=2, value="NĂM TAB CÒN LẠI").font = Font(bold=True, size=11, color="1F4E79")
    h += 1
    dat_header(ws, h, ["", "Tab", "Tầng", "Nội dung"])
    for ten, tang, mo in [
        ("TongHop", "🤖+🧑", "Bảng đầu bài: mỗi bậc prompt một dòng, số người và số máy cạnh nhau, kèm ngưỡng đánh giá."),
        ("TheoTin", "🧑", "10 CV × 4 bậc dạng bảng ngang — xem CV nào kéo tụt chỉ số."),
        ("NguoiCham_TungCau", "🧑", "Nguồn: out/<ver>/nguoi_cham_tung_cau.csv · mỗi câu tóm tắt một nhãn + lý do. Lọc được theo bậc, nhãn, cột 'Cần duyệt'."),
        ("NguoiCham_BoSot", "🧑", "Nguồn: out/<ver>/nguoi_cham_bo_sot.csv · mốc một bản tóm tắt tốt PHẢI nhắc mà AI quên — mẫu số của recall."),
        ("NguoiCham_TongKet", "🧑", "Nguồn: out/<ver>/nguoi_cham_tong_ket.csv · P/R/F1 từng CV + dòng TỔNG của mỗi bậc, kèm phân rã 4 kiểu lỗi."),
    ]:
        h += 1
        for i, g in enumerate(["", ten, tang, mo], start=1):
            if g:
                c = o(ws, h, i, g, wrap=True)
                if i == 2:
                    c.font = Font(bold=True)
        ws.row_dimensions[h].height = 30

    h += 2
    ws.cell(row=h, column=2, value="NĂM MÃ NHÃN (ranh giới khi phân vân: LUAT_NGUOI_CHAM.md)").font = Font(bold=True, size=11, color="1F4E79")
    for ma, mo in NHAN_MO_TA.items():
        h += 1
        o(ws, h, 2, ma).font = Font(bold=True)
        o(ws, h, 3, "TP" if ma == "DUNG" else "FP", ngang="center")
        o(ws, h, 4, mo)

    h += 2
    ws.cell(row=h, column=2, value="MUỐN CHẤM LẠI / SỬA NHÃN").font = Font(bold=True, size=11, color="1F4E79")
    for dong in [
        "Nhãn nằm trong 2_nguoi_cham_dien_nhan.py (bảng NHAN + BO_SOT) — KHÔNG sửa tay vào out/<ver>/*.csv (chạy lại script 1 là mất).",
        "Sửa xong chạy:  2_nguoi_cham_dien_nhan.py  →  3_nguoi_cham_tinh_diem.py --all  →  4_gop_ket_qua_excel.py",
        "Chấm theo hướng CHẶT: mọi số model tự tính là SAISO kể cả khi tính đúng; mốc thời gian bỏ sót phải là mốc của chính kinh nghiệm đó.",
    ]:
        h += 1
        ws.cell(row=h, column=2, value="• " + dong)
        ws.merge_cells(start_row=h, start_column=2, end_row=h, end_column=4)

    h += 2
    ws.cell(row=h, column=2, value="HẠN CHẾ PHẢI GHI TRONG BÁO CÁO").font = Font(bold=True, size=11, color="1F4E79")
    for dong in [
        "CV do người làm đề tài soạn, không phải hồ sơ thật (dữ liệu cá nhân + cần cài ca bẫy có đáp án).",
        "10 CV là ít; một bộ nhãn.",
        "TP đếm theo CÂU còn FN đếm theo MỐC → recall không cùng đơn vị, chỉ dùng để SO GIỮA CÁC BẬC.",
        "Thời gian đo trên GPU (RTX 5060), không phải CPU như máy demo.",
    ]:
        h += 1
        ws.cell(row=h, column=2, value="• " + dong)
        ws.merge_cells(start_row=h, start_column=2, end_row=h, end_column=4)

    h += 2
    c = ws.cell(row=h, column=2, value="⚠ " + NGUON_NHAN)
    c.fill, c.alignment = VANG, Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=h, start_column=2, end_row=h, end_column=4)
    ws.row_dimensions[h].height = 34


# ---------------------------------------------------------------- tab 2
def tab_tong_hop(wb, may, nguoi):
    ws = wb.create_sheet("TongHop")
    rong(ws, [8, 58, 9, 9, 8, 10, 11, 10, 10, 12, 12, 11, 13, 10, 10])

    tieu_de_to(ws, 1, "AI TÓM TẮT CV — TỔNG HỢP 4 BẬC PROMPT (ablation)", 15)
    ws.cell(row=2, column=1,
            value="10 CV đa ngành (7 ca thường + 3 ca bẫy) × 4 phiên bản · qwen3:8b qua Ollama · temperature=0 · chấm lượt 1")
    ws.cell(row=3, column=1, value="Tầng người = chấm theo câu (LUAT_NGUOI_CHAM.md) · Tầng máy = may_cham.py · " + NGUON_NHAN)

    dat_header(ws, 5, [
        "Bản", "Thêm gì so với bậc dưới", "Số câu", "DUNG", "Lỗi", "Mốc bỏ sót",
        "Precision", "Recall", "F1", "BIA+SAISO", "SAORONG",
        "CV có số lạ (máy)", "Đúng khuôn 3-5 câu", "Ổn định", "Giây/CV",
    ])
    ws.row_dimensions[5].height = 34

    h = 6
    for v in VERS:
        m = next((r for r in may if r["ver"].lower() == v), {})
        t = next((r for r in nguoi.get(v, []) if r["id"] == "TONG"), None)
        loi = sum(int(t[k]) for k in ("bia", "suydien", "saiso", "saorong")) if t else None
        so_cau = int(t["so_cau"]) if t else None
        gia_tri = [
            v.upper(), MO_TA_VER[v],
            so_cau, int(t["dung"]) if t else None, loi, int(t["bo_sot"]) if t else None,
            so(t["precision"]) if t else None, so(t["recall"]) if t else None, so(t["f1"]) if t else None,
            (int(t["bia"]) + int(t["saiso"])) / so_cau if t and so_cau else None,
            int(t["saorong"]) / so_cau if t and so_cau else None,
            f'{m.get("tin_co_so_la", "")}/10', so(m.get("dung_khuon_3_5_cau_%")) / 100 if m else None,
            so(m.get("on_dinh_tb")), so(m.get("giay_tb")),
        ]
        for i, g in enumerate(gia_tri, start=1):
            c = o(ws, h, i, g, wrap=(i == 2), ngang="left" if i == 2 else "center")
            if i in (7, 8, 9, 14):
                c.number_format = "0.000"
            elif i in (10, 11, 13):
                c.number_format = "0.0%"
            elif i == 15:
                c.number_format = "0.0"
            if v == "v4":
                c.fill = LUC
                c.font = Font(bold=True)
        ws.row_dimensions[h].height = 32
        h += 1

    t4 = next((r for r in nguoi.get("v4", []) if r["id"] == "TONG"), None)
    h += 1
    ws.cell(row=h, column=1, value="Ngưỡng đánh giá (chốt TRƯỚC khi đo — AI_TESTING_REFERENCE.md, xem LUAT_NGUOI_CHAM.md mục 4)").font = Font(bold=True)
    h += 1
    dat_header(ws, h, ["", "Chỉ số", "Tốt", "Chấp nhận được", "Cần cải thiện", "V4 đạt mức"])
    if t4:
        n4 = int(t4["so_cau"])
        chi_so = [
            ("Precision", so(t4["precision"]), "≥ 0.85", "0.70 – 0.84", "< 0.70", muc),
            ("Recall", so(t4["recall"]), "≥ 0.85", "0.70 – 0.84", "< 0.70", muc),
            ("F1", so(t4["f1"]), "≥ 0.85", "0.70 – 0.84", "< 0.70", muc),
            ("Tỉ lệ BIA + SAISO", (int(t4["bia"]) + int(t4["saiso"])) / n4, "0%", "≤ 5%", "> 5%",
             lambda x: "Tốt" if x == 0 else "Chấp nhận được" if x <= 0.05 else "Cần cải thiện"),
            ("Tỉ lệ SAORONG", int(t4["saorong"]) / n4, "≤ 5%", "≤ 15%", "> 15%",
             lambda x: "Tốt" if x <= 0.05 else "Chấp nhận được" if x <= 0.15 else "Cần cải thiện"),
        ]
        for ten, val, tot, cn, cc, ham in chi_so:
            h += 1
            kq = ham(val)
            hien = f"{val:.3f}" if ten in ("Precision", "Recall", "F1") else f"{val:.1%}"
            for i, g in enumerate(["", ten, tot, cn, cc, f"{hien} — {kq}"], start=1):
                if g:
                    c = o(ws, h, i, g)
                    if i == 6:
                        c.fill = LUC if kq == "Tốt" else VANG if kq == "Chấp nhận được" else DO

    tong = {v: next((r for r in nguoi.get(v, []) if r["id"] == "TONG"), None) for v in VERS}
    if all(tong.values()):
        def f(v, k):
            return so(tong[v][k])

        h += 2
        for dong in [
            "ĐỌC BẢNG NÀY THẾ NÀO:",
            f"• V1→V2 (ép JSON schema): precision TỤT {f('v1','precision'):.3f}→{f('v2','precision'):.3f}. Model viết lan tới trần "
            f"độ dài (3,4→9,4 câu/CV), đẻ {tong['v2']['saorong']} câu sáo rỗng — C08 lặp y hệt một câu 13 lần, C09 lặp 11 lần. "
            "Ép khuôn ép được HÌNH DẠNG, không thay được luật viết.",
            f"• V2→V3 (luật chống bịa): SAORONG {tong['v2']['saorong']}→{tong['v3']['saorong']}, precision nhảy lên {f('v3','precision'):.3f}. "
            f"Nhưng lỗi tự cộng số năm vẫn còn {tong['v3']['saiso']} câu — model không coi việc đó là bịa.",
            f"• V3→V4 (luật riêng của tóm tắt): SAISO {tong['v3']['saiso']}→{tong['v4']['saiso']}, precision giữ nguyên "
            f"({f('v3','precision'):.3f}→{f('v4','precision'):.3f}). Đổi lại recall tụt {f('v3','recall'):.3f}→{f('v4','recall'):.3f}: "
            "trần 3-5 câu bắt cắt bớt, và C09 mất trọn giai đoạn 2016–2020.",
            f"• Nói thẳng: theo F1, V4 ({f('v4','f1'):.3f}) THẤP HƠN V3 ({f('v3','f1'):.3f}). Cái V4 mua được là ít số tự tính hơn "
            "và bản tóm tắt ngắn đọc được — đánh đổi bằng vài mốc bị bỏ sót. Máy đo đã thấy trước (bỏ sót mốc 37,8% → 51,4%).",
            "• Precision đắt hơn recall ở bài này: mốc bị bỏ thì người tuyển dụng mở CV ra là thấy; con số sai thì họ tin luôn.",
        ]:
            c = ws.cell(row=h, column=1, value=dong)
            ws.merge_cells(start_row=h, start_column=1, end_row=h, end_column=15)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[h].height = 15 if dong.endswith(":") else 32
            h += 1

    ws.freeze_panes = "A6"


# ---------------------------------------------------------------- tab 3
def tab_theo_tin(wb, nguoi):
    ws = wb.create_sheet("TheoTin")
    rong(ws, [26, 36] + [8, 8, 7, 10, 10, 9] * 4)

    tieu_de_to(ws, 1, "TỪNG CV QUA 4 BẬC — CV nào kéo tụt chỉ số", 26)

    hang = 3
    for c, t in ((1, "Mã CV"), (2, "Ca")):
        cell = ws.cell(row=hang, column=c, value=t)
        cell.fill, cell.font = XANH, TRANG
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=hang, start_column=c, end_row=hang + 1, end_column=c)
    for k, v in enumerate(VERS):
        c0 = 3 + k * 6
        c = ws.cell(row=hang, column=c0, value=v.upper())
        c.fill, c.font = XANH, TRANG
        c.alignment = Alignment(horizontal="center")
        ws.merge_cells(start_row=hang, start_column=c0, end_row=hang, end_column=c0 + 5)
    dat_header(ws, hang + 1, ["", ""] + ["Số câu", "DUNG", "Sót", "Precision", "Recall", "F1"] * 4)
    ws.row_dimensions[hang + 1].height = 28

    h = hang + 2
    for cv in sorted(TEN_CV):
        o(ws, h, 1, cv)
        o(ws, h, 2, TEN_CV[cv], wrap=True)
        if cv in ("C01_khong_ghi_so_nam", "C09_khoang_trong_su_nghiep", "C10_liet_ke_cong_nghe"):
            for cc in (1, 2):
                ws.cell(row=h, column=cc).fill = VANG
        for k, v in enumerate(VERS):
            r = next((x for x in nguoi.get(v, []) if x["id"] == cv), None)
            if r is None:
                continue
            c0 = 3 + k * 6
            gia_tri = [int(r["so_cau"]), int(r["dung"]), int(r["bo_sot"]),
                       so(r["precision"]), so(r["recall"]), so(r["f1"])]
            for i, g in enumerate(gia_tri):
                cell = o(ws, h, c0 + i, g, ngang="center")
                if i >= 3:
                    cell.number_format = "0.000"
                    if i == 5:
                        cell.fill = to_mau(g)
        ws.row_dimensions[h].height = 26
        h += 1

    h += 1
    for dong in [
        "Tô vàng = ba ca bẫy về SỐ và SUY DIỄN (C01 không ghi số năm · C09 hai mốc rời · C10 liệt kê công nghệ).",
        "C09 là ca yếu nhất ở MỌI bậc: bậc nào cũng tự cộng thành '4 năm kinh nghiệm'; V4 còn bỏ mất cả giai đoạn 2016–2020.",
        "C10 ở V4: đúng cái bẫy đặt ra — biến dòng liệt kê Docker/Kubernetes thành 'có kinh nghiệm với…', 'mạnh nhất là triển khai hệ thống'.",
        "C08 ở V2: 18 câu, trong đó 13 câu y hệt nhau 'đạt yêu cầu về kinh nghiệm và kỹ năng…' — schema cho phép 1500 ký tự và model viết cho đủ.",
    ]:
        ws.cell(row=h, column=1, value=dong)
        h += 1

    ws.freeze_panes = "C5"


# ---------------------------------------------------------------- tab 4
def tab_tung_cau(wb, tung_cau):
    ws = wb.create_sheet("NguoiCham_TungCau")
    rong(ws, [7, 26, 8, 80, 11, 10, 60])

    tieu_de_to(ws, 1, "PHIẾU CHẤM TAY — MỖI CÂU TÓM TẮT MỘT DÒNG  (nguồn: out/<ver>/nguoi_cham_tung_cau.csv)", 7)
    ws.cell(row=2, column=1,
            value="Lọc cột 'Bậc' để xem riêng một phiên bản · lọc 'Nhãn' để xem một kiểu lỗi · lọc 'Cần duyệt' = ✓ để xem các dòng người duyệt phải đọc kỹ. "
                  "Sửa nhãn trong 2_nguoi_cham_dien_nhan.py rồi chạy lại, đừng sửa ở đây.")
    dat_header(ws, 4, ["Bậc", "Mã CV", "STT câu", "Câu AI viết", "Nhãn", "Cần duyệt", "Ghi chú / lý do chấm"])
    ws.row_dimensions[4].height = 26

    h = 5
    for ver, r in tung_cau:
        gc = r.get("ghi_chu", "") or ""
        can = "✓" if gc.startswith("PHÂN VÂN") else ""
        for i, g in enumerate([ver.upper(), r["id_tin"], int(r["stt_cau"]), r["cau"], r.get("nhan", ""), can, gc], start=1):
            c = o(ws, h, i, g, wrap=i in (4, 7), ngang="center" if i in (1, 3, 5, 6) else "left")
            if i == 5 and g:
                c.fill = LUC if g == "DUNG" else DO
                c.font = Font(bold=True)
            if can and i in (6, 7):
                c.fill = VANG
        h += 1

    ws.auto_filter.ref = f"A4:G{h - 1}"
    ws.freeze_panes = "A5"


# ---------------------------------------------------------------- tab 5
def tab_bo_sot(wb, bo_sot):
    ws = wb.create_sheet("NguoiCham_BoSot")
    rong(ws, [7, 26, 36, 11, 11, 80])

    tieu_de_to(ws, 1, "BỎ SÓT — MỐC MỘT BẢN TÓM TẮT TỐT PHẢI NHẮC MÀ AI QUÊN  (nguồn: out/<ver>/nguoi_cham_bo_sot.csv)", 6)
    for i, dong in enumerate([
        "Tab bên cạnh chỉ chấm được thứ AI ĐÃ viết ra. Thứ nó quên thì không có câu nào để chấm — phải đối chiếu 'moc_phai_co' trong dataset.json.",
        "Recall = DUNG / (DUNG + mốc bỏ sót). Không có cột này thì recall luôn bằng 1, và một bản tóm tắt chỉ viết 1 câu chắc ăn đạt precision 1.0.",
        "Diễn đạt khác cùng ý vẫn tính là có nhắc. Mốc thời gian phải là mốc của chính kinh nghiệm đó — năm tốt nghiệp trùng số không tính thay.",
    ]):
        ws.cell(row=2 + i, column=1, value="• " + dong)

    dat_header(ws, 6, ["Bậc", "Mã CV", "Ca", "Mốc phải có", "Bỏ sót", "Sót cái gì"])
    ws.row_dimensions[6].height = 26

    h = 7
    for ver, r in bo_sot:
        n = int(r["so_moc_bi_bo_sot"]) if (r.get("so_moc_bi_bo_sot") or "").strip().isdigit() else None
        gc = r.get("ghi_chu", "") or ""
        for i, g in enumerate([ver.upper(), r["id_tin"], TEN_CV.get(r["id_tin"], ""),
                               int(r["so_moc_phai_co"]), n, gc], start=1):
            c = o(ws, h, i, g, wrap=(i == 6), ngang="center" if i in (1, 4, 5) else "left")
            if i == 5 and n:
                c.fill = DO if n >= 2 else VANG
                c.font = Font(bold=True)
            if i == 6 and gc.startswith("PHÂN VÂN"):
                c.fill = VANG
        h += 1

    ws.auto_filter.ref = f"A6:F{h - 1}"
    ws.freeze_panes = "A7"


# ---------------------------------------------------------------- tab 6
def tab_tong_ket(wb, tong_ket):
    ws = wb.create_sheet("NguoiCham_TongKet")
    rong(ws, [7, 26, 36, 8, 8, 8, 7, 10, 8, 10, 11, 11, 10])

    tieu_de_to(ws, 1, "TỔNG KẾT TẦNG NGƯỜI — P / R / F1 TỪNG CV  (nguồn: out/<ver>/nguoi_cham_tong_ket.csv)", 13)
    ws.cell(row=2, column=1, value="Dòng TỔNG của mỗi bậc là con số đem trích vào báo cáo. " + NGUON_NHAN)

    dat_header(ws, 4, ["Bậc", "Mã CV", "Ca", "Số câu", "DUNG", "Sót",
                       "BIA", "SUYDIEN", "SAISO", "SAORONG", "Precision", "Recall", "F1"])
    ws.row_dimensions[4].height = 30

    h = 5
    for ver, r in tong_ket:
        la_tong = r["id"] == "TONG"
        gia_tri = [ver.upper(), r["id"], "— TỔNG CẢ BẬC —" if la_tong else TEN_CV.get(r["id"], ""),
                   int(r["so_cau"]), int(r["dung"]), int(r["bo_sot"]),
                   int(r["bia"]), int(r["suydien"]), int(r["saiso"]), int(r["saorong"]),
                   so(r["precision"]), so(r["recall"]), so(r["f1"])]
        for i, g in enumerate(gia_tri, start=1):
            c = o(ws, h, i, g, ngang="left" if i in (2, 3) else "center")
            if i >= 11:
                c.number_format = "0.000"
                if i == 13:
                    c.fill = to_mau(g)
            if la_tong:
                c.font = Font(bold=True)
                if i <= 10:
                    c.fill = XAM
        h += 1

    ws.auto_filter.ref = f"A4:M{h - 1}"
    ws.freeze_panes = "D5"


def chan_cong_thuc(wb):
    """Chuỗi bắt đầu bằng '=' bị openpyxl ghi thành công thức và Excel báo file hỏng."""
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f":
                    c.data_type = "s"


def workbook_mot_bac(ver, tung_cau, bo_sot, tong_ket):
    """Ba tab chấm tay của MỘT bậc, đặt trong out/<ver>/. CSV vẫn giữ để git diff được nhãn."""
    wb = Workbook()
    wb.remove(wb.active)
    tab_tung_cau(wb, [r for r in tung_cau if r[0] == ver])
    tab_bo_sot(wb, [r for r in bo_sot if r[0] == ver])
    tab_tong_ket(wb, [r for r in tong_ket if r[0] == ver])
    chan_cong_thuc(wb)
    dich = OUT / ver / f"NGUOI_CHAM_{ver}.xlsx"
    wb.save(dich)
    return dich


def main() -> int:
    may = doc(OUT / "may_cham_4_ban.csv")
    nguoi = {v: doc(OUT / v / "nguoi_cham_tong_ket.csv") for v in VERS}
    tung_cau = [(v, r) for v in VERS for r in doc(OUT / v / "nguoi_cham_tung_cau.csv")]
    bo_sot = [(v, r) for v in VERS for r in doc(OUT / v / "nguoi_cham_bo_sot.csv")]
    tong_ket = [(v, r) for v in VERS for r in nguoi[v]]

    wb = Workbook()
    wb.remove(wb.active)
    tab_doc_truoc(wb)
    tab_tong_hop(wb, may, nguoi)
    tab_theo_tin(wb, nguoi)
    tab_tung_cau(wb, tung_cau)
    tab_bo_sot(wb, bo_sot)
    tab_tong_ket(wb, tong_ket)
    chan_cong_thuc(wb)

    dich = OUT / "KET_QUA_TONG_HOP.xlsx"
    wb.save(dich)
    print(f"Đã ghi: {dich}")
    print(f"  NguoiCham_TungCau — {len(tung_cau)} câu · NguoiCham_BoSot — {len(bo_sot)} dòng · "
          f"NguoiCham_TongKet — {len(tong_ket)} dòng")
    for v in VERS:
        print(f"  {workbook_mot_bac(v, tung_cau, bo_sot, tong_ket)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
